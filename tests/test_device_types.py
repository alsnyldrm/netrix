import base64
import csv
import io
import json

import pytest
from netrix import models as m
from netrix.inventory.csv_transfer import render
from netrix.inventory.identity import apply_identity
from netrix.inventory.xlsx_transfer import from_csv
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError

BASE = "/api/v1/device-types"


def add_type(c, name="Controller"):
    response = c.post(BASE, json={"name": name, "description": "Kablosuz ağ"})
    assert response.status_code == 201, response.text
    return response.json()


def test_type_crud_links_rename_filter_and_archive_delete_guard(admin_client, inventory, env):
    c = admin_client
    kind = add_type(c)
    assert len(kind["code"]) <= 20
    response = c.post(
        "/api/v1/devices",
        json={
            "name": "Controller-01",
            "management_ip": "192.0.2.55",
            "vrf_id": inventory["vrf"]["id"],
            "kind": kind["code"],
        },
    )
    assert response.status_code == 201, response.text
    device = response.json()
    assert device["kind_name"] == "Controller"
    with env() as db:
        db.add(m.Interface(device_id=device["id"], name="port1", is_managed=True))
        db.commit()
    renamed = c.put(
        BASE + "/" + kind["id"], json={"name": "Kablosuz Denetleyici", "description": "Yeni açıklama"}
    )
    assert renamed.status_code == 200 and renamed.json()["code"] == kind["code"]
    assert c.get("/api/v1/devices/" + device["id"]).json()["kind_name"] == "Kablosuz Denetleyici"
    for resource, field in (("devices", "kind"), ("interfaces", "device_kind")):
        found = c.get("/api/v1/inventory/" + resource, params={"q": "Denetleyici"}).json()
        assert found["total"] == 1
        for value in (kind["code"], "Kablosuz Denetleyici"):
            filters = json.dumps([{"field": field, "op": "eq", "value": value}])
            assert c.get("/api/v1/inventory/" + resource, params={"filters": filters}).json()["total"] == 1
        filters = json.dumps([{"field": field, "op": "ne", "value": "Kablosuz Denetleyici"}])
        assert c.get("/api/v1/inventory/" + resource, params={"filters": filters}).json()["total"] == 0
    assert c.delete(BASE + "/" + kind["id"]).status_code == 409
    c.delete("/api/v1/devices/" + device["id"])
    blocked = c.delete(BASE + "/" + kind["id"])
    assert blocked.status_code == 409 and "arşiv dahil" in blocked.text
    c.post("/api/v1/inventory/devices/bulk", json={"ids": [device["id"]], "action": "restore"})
    assert (
        c.post(
            "/api/v1/inventory/devices/bulk",
            json={"ids": [device["id"]], "action": "update", "changes": {"kind": "switch"}},
        ).status_code
        == 200
    )
    assert c.get("/api/v1/devices/" + device["id"]).json()["kind_name"] == "Switch"
    assert c.delete(BASE + "/" + kind["id"]).status_code == 204
    assert c.get(BASE + "/" + kind["id"]).status_code == 404
    with env() as db:
        actions = set(db.scalars(select(m.Audit.action).where(m.Audit.entity_id == kind["id"])))
        assert actions == {"create", "update", "delete"}


def test_type_validation_default_and_database_fk(admin_client, inventory, env):
    c = admin_client
    assert c.post(BASE, json={"name": "   "}).status_code == 422
    assert c.post(BASE, json={"name": "Fake", "code": "switch"}).status_code == 422
    add_type(c)
    assert c.post(BASE, json={"name": "controller"}).status_code == 409
    other = next(r for r in c.get(BASE).json()["items"] if r["code"] == "other")
    assert c.delete(BASE + "/" + other["id"]).status_code == 409
    assert c.put(BASE + "/" + other["id"], json={"name": "Sınıflandırılmamış"}).status_code == 200
    assert (
        c.post(
            "/api/v1/devices",
            json={
                "name": "Bad",
                "management_ip": "192.0.2.55",
                "vrf_id": inventory["vrf"]["id"],
                "kind": "missing",
            },
        ).status_code
        == 422
    )
    with env() as db:
        db.add(
            m.Device(name="Bad", management_ip="192.0.2.55", vrf_id=inventory["vrf"]["id"], kind="missing")
        )
        with pytest.raises(IntegrityError):
            db.commit()


def test_deleted_detected_type_preserves_managed_type(admin_client, inventory, env):
    kind = add_type(admin_client)
    switch = next(r for r in admin_client.get(BASE).json()["items"] if r["code"] == "switch")
    assert admin_client.delete(BASE + "/" + switch["id"]).status_code == 204
    with env() as db:
        device = m.Device(
            name="Managed", management_ip="192.0.2.55", vrf_id=inventory["vrf"]["id"], kind=kind["code"]
        )
        db.add(device)
        db.flush()
        fields = apply_identity(db, device, {"kind": "switch", "vendor": "Aruba"}, None, "ssh")
        db.commit()
        assert device.kind == kind["code"] and device.vendor == "Aruba" and "kind" not in fields


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_types_in_templates_import_export_and_deleted_stale_selection(
    admin_client, transfer_profiles, inventory, format
):
    c = admin_client
    kind = add_type(c)
    url = "/api/v1/inventory/devices/" + format
    book = load_workbook(io.BytesIO(c.get("/api/v1/inventory/devices/xlsx/template").content))
    choices = book.defined_names["netrix_kind"]
    values = [
        cell.value for sheet, cells in choices.destinations for row in book[sheet][cells] for cell in row
    ]
    assert "Controller" in values
    content = render(
        [
            {
                "name": "Imported",
                "management_ip": "192.0.2.55",
                "snmp_credential_id": "SNMP-Profil-01",
                "config_credential_id": "SSH-Profil-01",
                "ssh_trust_host_key": "Hayır",
                "site_id": "Site-01",
                "adapter": "ios",
                "lifecycle_status": "Aktif",
                "kind": "Controller",
            }
        ]
    )
    data = content if format == "csv" else base64.b64encode(from_csv(content)).decode()
    result = c.post(url + "/import", json={"content": data, "apply": True})
    assert result.status_code == 200 and result.json()["imported"] == 1, result.text
    device = c.get("/api/v1/devices").json()["items"][0]
    assert device["kind"] == kind["code"]
    exported = c.post(url + "/export", json={"ids": [device["id"]]})
    if format == "csv":
        row = list(csv.DictReader(io.StringIO(exported.content.decode("utf-8-sig"))))[0]
    else:
        rows = list(load_workbook(io.BytesIO(exported.content)).active.values)
        row = dict(zip(rows[0], rows[1]))
    assert row["Cihaz türü"] == "Controller"
    c.post(
        "/api/v1/inventory/devices/bulk",
        json={"ids": [device["id"]], "action": "update", "changes": {"kind": "other"}},
    )
    assert c.delete(BASE + "/" + kind["id"]).status_code == 204
    report = c.post(url + "/import", json={"content": data}).json()
    assert report["errors"] and "Cihaz türü" in report["errors"][0]["message"]


@pytest.mark.parametrize("role,create_status", [("viewer", 403), ("operator", 201)])
def test_type_permissions(admin_client, env, role, create_status):
    with env() as db:
        db.scalar(select(m.User).where(m.User.email == "owner@example.test")).role = role
        db.commit()
    assert admin_client.get(BASE).status_code == 200
    result = admin_client.post(BASE, json={"name": "Controller"})
    assert result.status_code == create_status
    row = admin_client.get(BASE).json()["items"][0]
    assert admin_client.delete(BASE + "/" + row["id"]).status_code == 403
