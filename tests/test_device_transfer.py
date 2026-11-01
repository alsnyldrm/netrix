import base64
import io
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from netrix import models as m
from netrix.inventory.csv_transfer import render as render_csv
from netrix.inventory.xlsx_transfer import from_csv
from netrix.worker import claim, execute
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

pytestmark = pytest.mark.usefixtures("transfer_profiles")


def render(rows):
    return render_csv(
        [
            {
                "snmp_credential_id": "SNMP-Profil-01",
                "config_credential_id": "SSH-Profil-01",
                "ssh_trust_host_key": False,
                "site_id": "Site-01",
                "lifecycle_status": "active",
                "adapter": "ios",
                **row,
            }
            for row in rows
        ]
    )


BASE = "/api/v1/inventory/devices"


def encoded(data):
    return base64.b64encode(data).decode()


def workbook(rows):
    book = Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_templates_mark_required_columns_and_accept_only_required_values(admin_client, format, env):
    response = admin_client.get(f"{BASE}/{format}/template")
    assert response.status_code == 200 and f".{format}" in response.headers["content-disposition"]
    if format == "xlsx":
        book = load_workbook(io.BytesIO(response.content))
        assert book.sheetnames == ["Cihazlar", "Açıklamalar", "Seçenekler"]
        assert book["Cihazlar"].freeze_panes == "A2"
        assert [cell.value for cell in book["Cihazlar"][1]][:8] == [
            "Cihaz adı *",
            "Yönetim IP adresi *",
            "SNMP profili *",
            "SSH / API profili *",
            "Anahtara güven *",
            "Lokasyon *",
            "Yaşam döngüsü *",
            "Konfigürasyon adaptörü *",
        ]
        assert book["Cihazlar"]["C1"].comment.text == "Zorunlu alan"
        assert book["Cihazlar"]["I1"].comment.text.startswith("İsteğe bağlı")
        content = encoded(response.content)
    else:
        assert '"Cihaz adı *","Yönetim IP adresi *","SNMP profili *"' in response.text
        assert '"VRF *"' not in response.text
        content = response.text
    preview = admin_client.post(
        f"{BASE}/{format}/import", json={"content": content, "test_after": True}
    ).json()
    assert preview["valid"] == 1 and preview["imported"] == 0, preview
    with env() as db:
        assert db.scalar(select(func.count()).select_from(m.Job)) == 0
    result = admin_client.post(f"{BASE}/{format}/import", json={"content": content, "apply": True}).json()
    assert result["imported"] == 1 and result["tests"] == [], result
    with env() as db:
        assert db.scalar(select(m.Audit).where(m.Audit.action == format + "_import")) is not None


def test_xlsx_minimal_headers_sparse_optional_cells_and_invalid_rows(admin_client, env):
    data = workbook(
        [
            [
                "SNMP profili *",
                "SSH / API profili *",
                "Anahtara güven *",
                "Cihaz adı *",
                "Yönetim IP adresi *",
                "Lokasyon *",
                "Yaşam döngüsü *",
                "Konfigürasyon adaptörü *",
                "Model",
            ],
            [
                "SNMP-Profil-01",
                "SSH-Profil-01",
                False,
                "İstanbul, cihaz",
                "192.0.2.2",
                "Site-01",
                "Aktif",
                "ios",
            ],
            [
                "SNMP-Profil-01",
                "SSH-Profil-01",
                True,
                "İkinci",
                "invalid-ip",
                "Site-01",
                "Aktif",
                "ios",
                "X1",
            ],
        ]
    )
    result = admin_client.post(f"{BASE}/xlsx/import", json={"content": encoded(data), "apply": True}).json()
    assert result["imported"] == 0 and result["errors"][0]["row"] == 3
    with env() as db:
        assert db.scalar(select(func.count()).select_from(m.Device)) == 0
    data = workbook(
        [
            [
                "snmp_credential_id",
                "config_credential_id",
                "ssh_trust_host_key",
                "name",
                "management_ip",
                "site_id",
                "lifecycle_status",
                "adapter",
            ],
            ["SNMP-Profil-01", "SSH-Profil-01", False, "Minimal", "192.0.2.2", "Site-01", "Aktif", "ios"],
        ]
    )
    result = admin_client.post(f"{BASE}/xlsx/import", json={"content": encoded(data), "apply": True}).json()
    assert result["imported"] == 1, result


def test_xlsx_export_selected_only_masks_secrets_and_roundtrips(admin_client, inventory, env):
    content = render(
        [
            {
                "name": "=SUM(1,2)",
                "management_ip": "192.0.2.1",
                "vrf_id": "Global",
                "snmp_credential_id": "Test SNMP",
            }
        ]
    )
    admin_client.post(f"{BASE}/csv/import", json={"content": content, "apply": True})
    with env() as db:
        device = db.scalar(select(m.Device))
        id = device.id
        db.add(m.Device(name="Not selected", management_ip="192.0.2.3", vrf_id=inventory["vrf"]["id"]))
        db.commit()
    response = admin_client.post(f"{BASE}/xlsx/export", json={"ids": [id]})
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    sheet = book["Cihazlar"]
    assert sheet.max_row == 2 and sheet["A2"].data_type == "s"
    rows = list(sheet.values)
    data = dict(zip(rows[0], rows[1]))
    assert data["SNMP kimlik doğrulama anahtarı"] == data["SNMP gizlilik anahtarı"] == "********"
    assert "fixture-auth" not in str(rows) and "fixture-priv" not in str(rows)
    sheet["B2"] = "192.0.2.4"
    output = io.BytesIO()
    book.save(output)
    result = admin_client.post(
        f"{BASE}/xlsx/import", json={"content": encoded(output.getvalue()), "apply": True}
    ).json()
    assert result["imported"] == 1, result
    with env() as db:
        assert db.scalar(select(m.Device).where(m.Device.management_ip == "192.0.2.4")).name == "=SUM(1,2)"


@pytest.mark.parametrize(
    "data",
    [
        b"not a workbook",
        workbook([["name", "management_ip", "vrf_id"], ["=1+1", "192.0.2.1", "Global"]]),
        workbook([["name", "management_ip", "vrf_id"]] + [["a", "192.0.2.1", "Global"]] * 501),
        workbook([["name", "management_ip", "vrf_id"] + ["extra"] * 30]),
    ],
)
def test_xlsx_rejects_invalid_formulas_and_size_limits(admin_client, data):
    assert admin_client.post(f"{BASE}/xlsx/import", json={"content": encoded(data)}).status_code == 422


def test_xlsx_rejects_compressed_bomb_plaintext_secrets_and_invalid_base64(admin_client):
    output = io.BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b" " * (8 * 1024 * 1024 + 1))
    assert (
        admin_client.post(f"{BASE}/xlsx/import", json={"content": encoded(output.getvalue())}).status_code
        == 422
    )
    assert admin_client.post(f"{BASE}/xlsx/import", json={"content": "%%%"}).status_code == 422
    content = from_csv(
        render(
            [
                {
                    "name": "Device",
                    "management_ip": "192.0.2.1",
                    "vrf_id": "Global",
                    "ssh_password": "PRIVATE-SECRET",
                }
            ]
        )
    )
    response = admin_client.post(f"{BASE}/xlsx/import", json={"content": encoded(content)})
    assert response.json()["errors"] and "PRIVATE-SECRET" not in response.text


@pytest.mark.parametrize("format", ["csv", "xlsx"])
async def test_import_and_test_queues_assigned_protocols_and_reports_real_worker_results(
    admin_client, inventory, env, monkeypatch, format
):
    ssh = admin_client.post(
        "/api/v1/credentials",
        json={
            "name": "Import SSH",
            "protocol": "ssh",
            "secret": {"username": "fixture", "password": "PASSWORD-SECRET"},
        },
    ).json()
    content = render(
        [
            {
                "name": "Two profiles",
                "management_ip": "192.0.2.1",
                "vrf_id": "Global",
                "snmp_credential_id": "Test SNMP",
                "config_credential_id": ssh["id"],
            },
            {
                "name": "Retired",
                "management_ip": "192.0.2.2",
                "vrf_id": "Global",
                "lifecycle_status": "retired",
            },
            {
                "name": "Planned",
                "management_ip": "192.0.2.3",
                "vrf_id": "Global",
                "snmp_credential_id": "Test SNMP",
                "lifecycle_status": "planned",
            },
        ]
    )
    if format == "xlsx":
        content = encoded(from_csv(content))
    calls = []

    async def probe(device, protocol, secret, collector):
        calls.append(protocol)
        if protocol == "ssh":
            assert secret["password"] == "PASSWORD-SECRET"
            raise TimeoutError("PASSWORD-SECRET")
        assert secret["auth_key"] == "fixture-auth"
        return "SNMP bağlantısı başarılı."

    monkeypatch.setattr("netrix.worker.probe_connection", probe)
    url = f"{BASE}/{format}/import"
    preview = admin_client.post(url, json={"content": content, "test_after": True}).json()
    assert preview["valid"] == 3 and calls == []
    response = admin_client.post(url, json={"content": content, "apply": True, "test_after": True})
    report = response.json()
    assert report["imported"] == 3, report
    tests = report["tests"]
    queued = [row for row in tests if row["job_id"]]
    assert {row["protocol"] for row in queued} == {"SNMPv3", "SSH"}
    assert len([row for row in tests if row["status"] == "skipped"]) == 2
    assert calls == []
    # Only connection-test jobs exercise the fake protocol probe here.
    with env() as db:
        for job in db.scalars(select(m.Job).where(m.Job.kind.in_(("identify", "backup")))):
            job.status, job.active_key = "cancelled", None
        db.commit()
    ids = [row["job_id"] for row in queued]
    status = admin_client.post(f"{BASE}/import-tests/status", json={"ids": ids})
    assert all(row["status"] == "queued" for row in status.json())
    for _ in ids:
        job_id = claim()
        assert job_id in ids
        await execute(job_id)
    status = admin_client.post(f"{BASE}/import-tests/status", json={"ids": ids})
    assert {row["status"] for row in status.json()} == {"succeeded", "failed"}
    assert sorted(calls) == ["snmpv3", "ssh"] and "PASSWORD-SECRET" not in status.text
    assert (
        admin_client.post(url, json={"content": content, "apply": True, "test_after": True}).json()[
            "imported"
        ]
        == 0
    )
    with env() as db:
        assert db.scalar(select(func.count()).select_from(m.Job).where(m.Job.kind.like("test_%"))) == 2


def test_import_test_jobs_rollback_with_devices(admin_client, inventory, env, monkeypatch):
    from netrix import services

    original = services.enqueue

    def fail_test(db, kind, *args, **kwargs):
        if kind.startswith("test_"):
            from fastapi import HTTPException

            raise HTTPException(409, "Concurrent conflict")
        return original(db, kind, *args, **kwargs)

    monkeypatch.setattr(services, "enqueue", fail_test)
    content = render(
        [
            {
                "name": "Device",
                "management_ip": "192.0.2.1",
                "vrf_id": "Global",
                "snmp_credential_id": "Test SNMP",
            }
        ]
    )
    result = admin_client.post(
        f"{BASE}/csv/import", json={"content": content, "apply": True, "test_after": True}
    )
    assert result.status_code == 409
    with env() as db:
        assert db.scalar(select(func.count()).select_from(m.Device)) == 0
        assert db.scalar(select(func.count()).select_from(m.Job)) == 0


def test_transfer_permissions_and_job_status_scope(admin_client, env):
    with env() as db:
        job = m.Job(kind="discovery", target_id="not-a-device")
        db.add(job)
        db.commit()
        id = job.id
    assert admin_client.post(f"{BASE}/import-tests/status", json={"ids": [id]}).status_code == 404
    admin_client.post(
        "/api/v1/users",
        json={
            "email": "viewer@transfer.test",
            "display_name": "Viewer",
            "password": "test-password-long",
            "role": "viewer",
        },
    )
    admin_client.post("/api/v1/auth/logout")
    admin_client.post(
        "/api/v1/auth/login", json={"email": "viewer@transfer.test", "password": "test-password-long"}
    )
    assert admin_client.get(f"{BASE}/xlsx/template").status_code == 403
    assert admin_client.post(f"{BASE}/xlsx/import", json={"content": "x"}).status_code == 403
    assert admin_client.post(f"{BASE}/xlsx/export", json={"ids": [id]}).status_code == 403
    assert admin_client.post(f"{BASE}/import-tests/status", json={"ids": [id]}).status_code == 403


@pytest.mark.parametrize("format", ["csv", "xlsx"])
@pytest.mark.parametrize(
    "field",
    [
        "name",
        "management_ip",
        "snmp_credential_id",
        "config_credential_id",
        "ssh_trust_host_key",
        "site_id",
        "lifecycle_status",
        "adapter",
    ],
)
@pytest.mark.parametrize("missing_column", [False, True])
def test_import_access_fields_are_required_before_any_writes(
    admin_client, env, format, field, missing_column
):
    import csv

    from netrix.inventory.csv_transfer import FIELDS

    row = {
        "snmp_credential_id": "SNMP-Profil-01",
        "config_credential_id": "SSH-Profil-01",
        "ssh_trust_host_key": "Hayır",
        "name": "Required fields",
        "management_ip": "192.0.2.150",
        "vrf_id": "Global",
        "site_id": "Site-01",
        "lifecycle_status": "Aktif",
        "adapter": "ios",
    }
    output = io.StringIO()
    if missing_column:
        del row[field]
    else:
        row[field] = "  "
    writer = csv.DictWriter(output, fieldnames=list(row))
    writer.writeheader()
    writer.writerow(row)
    content = output.getvalue() if format == "csv" else encoded(workbook([list(row), list(row.values())]))
    for apply in (False, True):
        response = admin_client.post(
            f"{BASE}/{format}/import", json={"content": content, "apply": apply, "test_after": True}
        )
        if missing_column:
            assert response.status_code == 422 and FIELDS[field] in response.json()["detail"]
        else:
            report = response.json()
            assert report["imported"] == 0 and report["valid"] == 0
            assert report["errors"] == [
                {"row": 2, "field": FIELDS[field], "message": "Zorunlu alan boş bırakılamaz"}
            ]
        with env() as db:
            assert db.scalar(select(func.count()).select_from(m.Device)) == 0
            assert db.scalar(select(func.count()).select_from(m.Job)) == 0


@pytest.mark.parametrize("format", ["csv", "xlsx"])
@pytest.mark.parametrize("value,expected", [("Hayır", False), ("Evet", True), (False, False), (True, True)])
def test_import_requires_explicit_trust_choice_and_previews_profiles(
    admin_client, env, format, value, expected
):
    content = render(
        [
            {
                "name": "Explicit trust",
                "management_ip": "192.0.2.151",
                "vrf_id": "Global",
                "ssh_trust_host_key": value,
            }
        ]
    )
    if format == "xlsx":
        content = encoded(from_csv(content))
    preview = admin_client.post(f"{BASE}/{format}/import", json={"content": content}).json()
    assert preview["valid"] == 1, preview
    assert preview["preview"][0]["snmp_profile"] == "SNMP-Profil-01"
    assert preview["preview"][0]["config_profile"] == "SSH-Profil-01"
    assert preview["preview"][0]["ssh_trust_host_key"] is expected
    result = admin_client.post(f"{BASE}/{format}/import", json={"content": content, "apply": True}).json()
    assert result["imported"] == 1, result
    with env() as db:
        assert db.scalar(select(m.Device)).ssh_trust_host_key is expected


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_invalid_trust_does_not_default_to_false(admin_client, format):
    content = render(
        [
            {
                "name": "Invalid trust",
                "management_ip": "192.0.2.152",
                "vrf_id": "Global",
                "ssh_trust_host_key": "maybe",
            }
        ]
    )
    if format == "xlsx":
        content = encoded(from_csv(content))
    report = admin_client.post(f"{BASE}/{format}/import", json={"content": content, "apply": True}).json()
    assert report["imported"] == 0 and report["errors"][0]["field"] == "Anahtara güven"
