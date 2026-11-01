import base64
import csv
import io
from zipfile import ZipFile

import pytest
from netrix import models as m
from netrix.inventory.csv_transfer import FIELDS, REQUIRED, SECRETS, TEMPLATE_FIELDS, render
from openpyxl import load_workbook
from sqlalchemy import event, select

pytestmark = pytest.mark.usefixtures("transfer_profiles")
BASE = "/api/v1/inventory/devices"


def values(book, key):
    title, bounds = next(book.defined_names["netrix_" + key].destinations)
    return [cell.value for row in book[title][bounds] for cell in row if cell.value is not None]


def test_xlsx_lists_are_current_protocol_filtered_public_and_literal(admin_client, env, institution):
    # Different protocols, formula-like names and duplicate locations exercise real references.
    for name, protocol, secret in (
        ("v2c-only", "snmpv2c", {"community": "DO-NOT-EXPORT-COMMUNITY"}),
        ("API-only", "rest", {"token": "DO-NOT-EXPORT-TOKEN"}),
        ('=HYPERLINK("fixture")', "ssh", {"username": "fixture", "password": "DO-NOT-EXPORT-PASSWORD"}),
    ):
        assert (
            admin_client.post(
                "/api/v1/credentials", json={"name": name, "protocol": protocol, "secret": secret}
            ).status_code
            == 201
        )
    with env() as db:
        a, b = m.Site(name="Same site", code="S1"), m.Site(name="Same site", code="S2")
        db.add_all([a, b])
        db.flush()
        duplicates = {a.id, b.id}
        vrf = db.scalar(select(m.VRF))
        db.add(
            m.Device(
                name="Existing",
                management_ip="192.0.2.100",
                vrf_id=vrf.id,
                vendor="Known vendor",
                model="Known model",
            )
        )
        db.add(m.BackupGroup(name="Switches"))
        db.commit()
        secrets = list(db.scalars(select(m.Credential.encrypted_secret)))
        statements = []
        engine = db.get_bind()

        def capture(connection, cursor, statement, parameters, context, many):
            statements.append(statement)

        event.listen(engine, "before_cursor_execute", capture)
        try:
            response = admin_client.get(BASE + "/xlsx/template")
        finally:
            event.remove(engine, "before_cursor_execute", capture)
    assert response.status_code == 200, response.text
    assert not any("encrypted_secret" in sql for sql in statements)
    with ZipFile(io.BytesIO(response.content)) as archive:
        xml = "\n".join(archive.read(name).decode() for name in archive.namelist() if name.endswith(".xml"))
    assert all(
        secret not in xml
        for secret in [*secrets, "DO-NOT-EXPORT-COMMUNITY", "DO-NOT-EXPORT-TOKEN", "DO-NOT-EXPORT-PASSWORD"]
    )
    book = load_workbook(io.BytesIO(response.content))
    assert values(book, "snmp_credential_id") == ["SNMP-Profil-01", "v2c-only"]
    assert set(values(book, "config_credential_id")) == {
        '\'=HYPERLINK("fixture")',
        "API-only",
        "SSH-Profil-01",
    }
    assert duplicates <= set(values(book, "site_id")) and "Same site" not in values(book, "site_id")
    assert values(book, "backup_group_id") == ["Switches"]
    assert values(book, "vendor") == ["Known vendor"]
    assert values(book, "model") == ["Known model"]
    assert values(book, "vrf_id") == ["Global"]
    assert values(book, "ssh_trust_host_key") == ["Hayır", "Evet"]
    assert "Aruba 61xx · AOS-CX" in values(book, "adapter")
    assert book["Seçenekler"].sheet_state == "hidden"
    assert not any(c.data_type in ("f", "e") or c.hyperlink for sheet in book for row in sheet for c in row)
    validations = book["Cihazlar"].data_validations.dataValidation
    assert len(validations) == len(FIELDS) + len(SECRETS)
    for validation, key in zip(validations, (*TEMPLATE_FIELDS, *SECRETS)):
        assert validation.showDropDown is False  # OOXML true hides the arrow.
        assert validation.allowBlank is (key not in REQUIRED)
        assert validation.type == "list" and validation.formula1 == "=netrix_" + key
        assert len(validation.cells.ranges) == 1
        cell_range = next(iter(validation.cells.ranges))
        assert cell_range.min_row == 2 and cell_range.max_row == 501
    # No cached template: later records must appear on the next download.
    admin_client.post(
        "/api/v1/sites", json={"name": "New location", "code": "NEW", "institution_id": institution["id"]}
    )
    fresh = load_workbook(io.BytesIO(admin_client.get(BASE + "/xlsx/template").content))
    assert "New location" in values(fresh, "site_id")


def test_more_than_500_choices_roundtrip_without_relaxing_device_row_limit(admin_client, env):
    with env() as db:
        db.add_all(m.Site(name=f"Location {i:03}", code=f"L{i}") for i in range(505))
        db.commit()
    response = admin_client.get(BASE + "/xlsx/template")
    book = load_workbook(io.BytesIO(response.content))
    assert len(values(book, "site_id")) == 506
    assert book["Seçenekler"].max_row == 507
    result = admin_client.post(
        BASE + "/xlsx/import", json={"content": base64.b64encode(response.content).decode()}
    )
    assert result.json()["valid"] == 1, result.text
    book["Cihazlar"]["A502"] = "Too many devices"
    out = io.BytesIO()
    book.save(out)
    assert (
        admin_client.post(
            BASE + "/xlsx/import", json={"content": base64.b64encode(out.getvalue()).decode()}
        ).status_code
        == 422
    )


@pytest.mark.parametrize("vrf", [None, "", "Other"])
def test_optional_vrf_defaults_to_global_but_preserves_explicit_selection(admin_client, env, vrf):
    other = admin_client.post("/api/v1/vrfs", json={"name": "Other"}).json()
    with env() as db:
        db.add(
            m.Scope(
                name="Other network",
                cidr="192.0.2.0/24",
                vrf_id=other["id"],
                vlan_vid=20,
                credential_id=db.scalar(select(m.Credential.id).where(m.Credential.protocol == "snmpv3")),
            )
        )
        db.commit()
    row = dict(
        name="Optional VRF",
        management_ip="192.0.2.150",
        snmp_credential_id="SNMP-Profil-01",
        config_credential_id="SSH-Profil-01",
        ssh_trust_host_key=False,
        site_id="Site-01",
        lifecycle_status="Aktif",
        adapter="Aruba 54xx · ArubaOS-Switch",
    )
    if vrf is not None:
        row["vrf_id"] = vrf
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=list(row))
    writer.writeheader()
    writer.writerow(row)
    preview = admin_client.post(BASE + "/csv/import", json={"content": output.getvalue()}).json()
    assert preview["valid"] == 1, preview
    assert preview["preview"][0]["vrf"] == (vrf or "Global")
    result = admin_client.post(
        BASE + "/csv/import", json={"content": output.getvalue(), "apply": True}
    ).json()
    assert result["imported"] == 1, result
    with env() as db:
        device = db.scalar(select(m.Device))
        assert db.get(m.VRF, device.vrf_id).name == (vrf or "Global")
        assert (device.vrf_id == other["id"]) is bool(vrf)
        assert device.adapter == "aruba_5400"


def test_missing_default_vrf_fails_without_creating_one(admin_client, env):
    with env() as db:
        db.scalar(select(m.VRF)).name = "Renamed"
        db.commit()
    content = render(
        [
            dict(
                name="No default",
                management_ip="192.0.2.9",
                snmp_credential_id="SNMP-Profil-01",
                config_credential_id="SSH-Profil-01",
                ssh_trust_host_key=False,
                site_id="Site-01",
                lifecycle_status="active",
                adapter="ios",
            )
        ]
    )
    result = admin_client.post(BASE + "/csv/import", json={"content": content, "apply": True}).json()
    assert result["imported"] == 0 and "Global VRF bulunamadı" in result["errors"][0]["message"]
    with env() as db:
        assert list(db.scalars(select(m.Device))) == []
        assert [v.name for v in db.scalars(select(m.VRF))] == ["Renamed"]


def test_removed_profile_is_revalidated_when_old_workbook_is_uploaded(admin_client, env):
    response = admin_client.get(BASE + "/xlsx/template")
    with env() as db:
        db.delete(db.scalar(select(m.Credential).where(m.Credential.name == "SSH-Profil-01")))
        db.commit()
    result = admin_client.post(
        BASE + "/xlsx/import", json={"content": base64.b64encode(response.content).decode(), "apply": True}
    ).json()
    assert result["imported"] == 0 and result["errors"]
