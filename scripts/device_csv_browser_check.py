"""CSV/XLSX browser workflows against the local fake-probe transfer fixture."""

import csv
import io
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from playwright.sync_api import expect, sync_playwright

BASE = "http://127.0.0.1:8215"
with sync_playwright() as p:
    browser = p.chromium.launch()
    page = browser.new_page(
        viewport={"width": 1440, "height": 1000},
        bypass_csp=True,
        reduced_motion="reduce",
        accept_downloads=True,
    )
    errors, reports = [], {}
    page.on("pageerror", lambda e: errors.append(str(e)))
    page.route_web_socket("**/api/v1/ws", lambda socket: None)
    page.goto(BASE + "/#devices")
    page.get_by_label("E-posta", exact=True).fill("browser@example.test")
    page.get_by_label("Parola", exact=True).fill("browser-fixture-password")
    page.get_by_role("button", name="Çalışma alanına gir", exact=True).click()
    page.wait_for_load_state("networkidle")
    export = page.get_by_role("button", name="Dışarı Aktar", exact=True)
    import_button = page.get_by_role("button", name="İçeri Aktar", exact=True)
    expect(export).to_be_disabled()
    page.get_by_role("checkbox", name="Fixture 34 seç", exact=True).check()
    page.get_by_role("checkbox", name="Fixture 33 seç", exact=True).check()
    import_button.focus()
    page.keyboard.press("ArrowDown")
    expect(page.get_by_role("menuitem", name="CSV", exact=True)).to_be_focused()
    page.keyboard.press("ArrowDown")
    expect(page.get_by_role("menuitem", name="XLSX", exact=True)).to_be_focused()
    page.keyboard.press("Escape")
    expect(import_button).to_be_focused()
    expect(page.get_by_role("menu")).to_have_count(0)

    def screen(name):
        for theme in ("light", "dark"):
            for width, height in ((1440, 1000), (375, 812)):
                page.set_viewport_size({"width": width, "height": height})
                page.evaluate("t=>document.documentElement.dataset.theme=t", theme)
                assert page.locator("body").evaluate("e=>e.scrollWidth<=innerWidth"), name
                if page.get_by_role("dialog").count():
                    dialog = page.get_by_role("dialog")
                    assert dialog.evaluate("e=>e.scrollWidth<=e.clientWidth"), name
                    summary = dialog.locator(".csv-summary")
                    if summary.count():
                        assert (
                            abs(
                                summary.bounding_box()["x"]
                                - dialog.get_by_role("heading").bounding_box()["x"]
                            )
                            < 1
                        )
                    buttons = dialog.locator(".dialog-footer button")
                    if buttons.count():
                        buttons.last.scroll_into_view_if_needed()
                        boxes = [buttons.nth(i).bounding_box() for i in range(buttons.count())]
                        assert (
                            max(box["height"] for box in boxes) - min(box["height"] for box in boxes) < 1
                        ), boxes
                page.add_script_tag(path=str(Path("frontend/node_modules/axe-core/axe.min.js").resolve()))
                result = page.evaluate(
                    "async()=>await axe.run(document,{runOnly:{type:'tag',values:['wcag2a','wcag2aa','wcag21aa','wcag22aa']}})"
                )
                reports[f"{name}-{theme}-{width}"] = [v["id"] for v in result["violations"]]
                page.screenshot(path=f"artifacts/device-transfer-{name}-{theme}-{width}.png", full_page=True)
        page.set_viewport_size({"width": 1440, "height": 1000})

    import_button.click()
    screen("menu")
    page.keyboard.press("Escape")
    for format in ("csv", "xlsx"):
        export.click()
        with page.expect_download() as download:
            page.get_by_role("menuitem", name=format.upper(), exact=True).click()
        payload = Path(download.value.path()).read_bytes()
        assert download.value.suggested_filename.endswith("." + format)
        if format == "csv":
            rows = list(csv.DictReader(io.StringIO(payload.decode("utf-8-sig"))))
        else:
            values = list(load_workbook(io.BytesIO(payload)).active.values)
            rows = [dict(zip(values[0], row)) for row in values[1:]]
        assert {row["Cihaz adı"] for row in rows} == {"Fixture 33", "Fixture 34"}
        assert all(row["SSH parolası"] == "********" for row in rows)
        assert "fixture-password" not in str(rows)

    for format, start in (("csv", 100), ("xlsx", 110)):
        import_button.click()
        page.get_by_role("menuitem", name=format.upper(), exact=True).click()
        dialog = page.get_by_role("dialog")
        expect(dialog).to_have_accessible_name(f"Cihazları {format.upper()} ile içeri aktar")
        with page.expect_download() as template:
            page.get_by_role("button", name=f"{format.upper()} şablonu indir", exact=True).click()
        data = Path(template.value.path()).read_bytes()
        assert template.value.suggested_filename.endswith("." + format)
        if format == "xlsx":
            assert load_workbook(io.BytesIO(data)).active["A1"].value == "Cihaz adı *"
        else:
            assert "Cihaz adı *" in data.decode("utf-8-sig")
        file_input = page.get_by_label(f"{format.upper()} dosyası", exact=True)
        file_input.set_input_files(
            {"name": "bad." + format, "mimeType": "application/octet-stream", "buffer": b"invalid"}
        )
        expect(dialog.get_by_role("alert")).to_be_visible()
        expect(dialog.get_by_role("button", name="İçeri Aktar ve Test Et", exact=True)).to_be_disabled()
        rows = [
            ["Cihaz adı *", "Yönetim IP adresi *", "VRF *", "SSH / API profili"],
            [format + " başarılı", f"192.0.2.{start + 1}", "Global", "Fixture SSH"],
            [format + " başarısız", f"192.0.2.{start + 2}", "Global", "Fixture SSH"],
            [format + " profilsiz", f"192.0.2.{start + 3}", "Global", ""],
        ]
        if format == "csv":
            output = io.StringIO()
            csv.writer(output, delimiter=";").writerows(rows)
            payload = output.getvalue().encode()
        else:
            book = Workbook()
            for row in rows:
                book.active.append(row)
            output = io.BytesIO()
            book.save(output)
            payload = output.getvalue()
        file_input.set_input_files(
            {"name": "devices." + format, "mimeType": "application/octet-stream", "buffer": payload}
        )
        expect(dialog.locator(".csv-summary")).to_contain_text("3 geçerli")
        screen(format + "-preview")
        # Exercise the import-and-test action on a phone, including footer reachability.
        page.set_viewport_size({"width": 375, "height": 812})
        dialog.get_by_role("button", name="İçeri Aktar ve Test Et", exact=True).click()
        expect(dialog).to_have_accessible_name("İçe aktarma · Bağlantı testi sonuçları")
        expect(dialog.locator(".csv-summary")).to_contain_text("2 bekleyen · 1 test edilemedi")
        page.request.post(BASE + "/_fixture/complete-tests", headers={"X-Netrix-Request": "1"})
        expect(dialog.locator(".csv-summary")).to_have_text(
            "1 başarılı · 1 başarısız · 0 bekleyen · 1 test edilemedi", timeout=15000
        )
        expect(dialog).not_to_contain_text("fixture-password")
        screen(format + "-results")
        dialog.get_by_role("button", name="Kapat", exact=True).click()
        # Re-import rejects duplicates without creating more tests.
        import_button.click()
        page.get_by_role("menuitem", name=format.upper(), exact=True).click()
        file_input.set_input_files(
            {"name": "devices." + format, "mimeType": "application/octet-stream", "buffer": payload}
        )
        expect(dialog.locator(".csv-summary")).to_contain_text("3 hatalı satır")
        expect(dialog.get_by_role("button", name="İçeri Aktar", exact=True)).to_be_disabled()
        page.keyboard.press("Escape")
    # Plain import remains available without connection tests.
    import_button.click()
    page.get_by_role("menuitem", name="CSV", exact=True).click()
    page.get_by_label("CSV dosyası", exact=True).set_input_files(
        {
            "name": "minimal.csv",
            "mimeType": "text/csv",
            "buffer": b"name,management_ip,vrf_id\nPlain import,192.0.2.120,Global",
        }
    )
    dialog = page.get_by_role("dialog")
    expect(dialog.locator(".csv-summary")).to_contain_text("1 geçerli")
    dialog.get_by_role("button", name="İçeri Aktar", exact=True).click()
    expect(dialog.locator(".csv-summary")).to_have_text("1 cihaz içeri aktarıldı.")
    assert not any(reports.values()), reports
    assert not errors, errors
    print(
        json.dumps(
            {
                "csv_xlsx_workflow": "PASS",
                "selected_only": True,
                "masked_secrets": True,
                "worker_test_results": "1 success, 1 failure, 1 skipped per format",
                "axe": reports,
                "console_errors": errors,
            }
        )
    )
    browser.close()
